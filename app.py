# -*- coding: utf-8 -*-
"""
ホテル収支計画 (Proforma) 作成アプリ — Flask ローカル Web アプリ

使い方:
    pip install -r requirements.txt
    python app.py
    ブラウザで http://127.0.0.1:5000 を開く

ブラウザ上のフォームに前提値を日本語で入力 →「Excelを生成」を押すと、
元のサンプル(template.xlsx)の数式・レイアウトをすべて保持したまま、
入力した前提値だけを埋め込んだ完成版 Excel がダウンロードされます。
Excel で開くと自動で再計算され、精緻なモデルが完成します。
"""
import io
import os
from copy import copy

from flask import Flask, render_template, request, send_file
import openpyxl

import fields as F
import translations as T

APP_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(APP_DIR, "template.xlsx")

app = Flask(__name__)


# --------------------------------------------------------------------------- #
# テンプレートからの既定値読み取り
# --------------------------------------------------------------------------- #
def _raw(ws, cell):
    v = ws[cell].value
    if isinstance(v, str) and v.startswith("="):
        return None  # 数式セルは既定値なし
    return v


def _disp(kind, v):
    """セルの生値を画面表示用の値に変換。"""
    if v is None:
        return "" if kind == "text" else 0
    if kind == "pct":
        return round(float(v) * 100, 6)
    if kind == "int":
        try:
            return int(round(float(v)))
        except (TypeError, ValueError):
            return v
    return v


def build_context():
    """テンプレートの既定値を読み込み、フォーム描画用のデータ構造を返す。"""
    wb = openpyxl.load_workbook(TEMPLATE_PATH, data_only=False)

    sections = []
    for sec in F.SECTIONS:
        fl = []
        for fd in sec["fields"]:
            ws = wb[fd["sheet"]]
            fl.append({
                **fd,
                "name": f"f__{fd['sheet']}__{fd['cell']}",
                "value": _disp(fd["kind"], _raw(ws, fd["cell"])),
            })
        sections.append({"id": sec["id"], "title": sec["title"], "fields": fl})

    # 客室タイプ
    rt_ws = wb["Room Types"]
    rooms_ws = wb["Rooms"]
    room_types = []
    for rt in F.room_type_fields():
        i = rt["idx"]
        name = _raw(rt_ws, rt["name_cell"][1])
        rooms = _raw(rt_ws, rt["rooms_cell"][1])
        size = _raw(rt_ws, rt["size_cells"][0][1])
        occ = _raw(rooms_ws, rt["occ_cell"][1])
        rate = _raw(rooms_ws, rt["rate_cell"][1])
        room_types.append({
            "idx": i,
            "name": name or "",
            "rooms": rooms or 0,
            "size": size or 0,
            "occ": round(float(occ) * 100, 4) if occ not in (None, "") else 0,
            "rate": rate or 0,
        })

    # 人員計画
    staffing = []
    for grp in F.STAFFING:
        ws = wb[grp["sheet"]]
        roles = []
        for r in grp["roles"]:
            roles.append({
                "label": r["label"],
                "salary_name": f"s__{grp['sheet']}__{r['salary_cell']}",
                "head_name": f"h__{grp['sheet']}__{r['head_cell']}",
                "salary": _raw(ws, r["salary_cell"]) or 0,
                "head": _raw(ws, r["head_cell"]) or 0,
            })
        staffing.append({"id": grp["id"], "title": grp["title"], "roles": roles})

    return sections, room_types, staffing


# --------------------------------------------------------------------------- #
# ルーティング
# --------------------------------------------------------------------------- #
@app.route("/")
def index():
    sections, room_types, staffing = build_context()
    return render_template("index.html",
                           sections=sections,
                           room_types=room_types,
                           staffing=staffing)


def _to_num(s, kind):
    s = (s or "").strip().replace(",", "")
    if s == "":
        return None
    try:
        v = float(s)
    except ValueError:
        return None
    if kind == "pct":
        return v / 100.0
    if kind == "int":
        return int(round(v))
    return v


@app.route("/generate", methods=["POST"])
def generate():
    wb = openpyxl.load_workbook(TEMPLATE_PATH, data_only=False)
    form = request.form

    # 平坦フィールド
    for sec in F.SECTIONS:
        for fd in sec["fields"]:
            key = f"f__{fd['sheet']}__{fd['cell']}"
            if key not in form:
                continue
            ws = wb[fd["sheet"]]
            if fd["kind"] == "text":
                val = form.get(key, "").strip()
                ws[fd["cell"]] = val
            else:
                val = _to_num(form.get(key), fd["kind"])
                if val is None:
                    continue
                ws[fd["cell"]] = val
                for extra in fd.get("spread", []):
                    ws[extra] = val

    # 客室タイプ
    rt_ws = wb["Room Types"]
    rooms_ws = wb["Rooms"]
    for rt in F.room_type_fields():
        i = rt["idx"]
        name = form.get(f"rt_name_{i}", "").strip()
        rt_ws[rt["name_cell"][1]] = name if name else None
        rooms = _to_num(form.get(f"rt_rooms_{i}"), "int")
        if rooms is not None:
            rt_ws[rt["rooms_cell"][1]] = rooms
        size = _to_num(form.get(f"rt_size_{i}"), "num")
        if size is not None:
            for (_, c) in rt["size_cells"]:
                rt_ws[c] = size
        occ = _to_num(form.get(f"rt_occ_{i}"), "pct")
        if occ is not None:
            rooms_ws[rt["occ_cell"][1]] = occ
        rate = _to_num(form.get(f"rt_rate_{i}"), "num")
        if rate is not None:
            rooms_ws[rt["rate_cell"][1]] = rate

    # 人員計画
    for grp in F.STAFFING:
        ws = wb[grp["sheet"]]
        for r in grp["roles"]:
            sal = _to_num(form.get(f"s__{grp['sheet']}__{r['salary_cell']}"), "num")
            if sal is not None:
                ws[r["salary_cell"]] = sal
            head = _to_num(form.get(f"h__{grp['sheet']}__{r['head_cell']}"), "int")
            if head is not None:
                ws[r["head_cell"]] = head

    # 出力言語の適用（日本語 / 英語）
    lang = form.get("lang", "ja")
    T.translate_workbook(wb, lang)

    # Excel で開いたときに全シートを再計算させる
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=("hotel_proforma_ja.xlsx" if lang == "ja" else "hotel_proforma_en.xlsx"),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)
